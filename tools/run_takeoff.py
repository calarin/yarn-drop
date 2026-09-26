#!/usr/bin/env python3
"""Fourth-floor furniture & equipment plan takeoff (A06.04A east / A06.04B west).

    python tools/run_takeoff.py --east inputs/4th_east.pdf --west inputs/4th_west.pdf --out deliverables

Writes the room x code table, instance list, callout catalog, exceptions, component-tag list,
overlays (PNG + PDF), pilot-room close-ups and the reconciliation summary. The source PDFs
and any model are only read.
"""
import argparse, csv, json, os, sys
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, 'takeoff'))
sys.path.insert(0, HERE)

def write_csv(path, rows, fields=None):
    fields = fields or (list(rows[0].keys()) if rows else [])
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore'); w.writeheader()
        for r in rows:
            w.writerow({k: (', '.join(map(str, v)) if isinstance(v, (list, tuple)) else v) for k, v in r.items() if k in fields})

ROW_FIELDS = ['row_id', 'sheet', 'room_number', 'room_name', 'drawing_code', 'code_description', 'item_class', 'unit',
              'plan_qty', 'model_qty', 'difference', 'model_status', 'source_sheet', 'callout_ids', 'typ_on_label',
              'target_symbol', 'counted_ids', 'evidence', 'confidence', 'review_status', 'pilot']
INST_FIELDS = ['instance_id', 'sheet', 'room_number', 'room_name', 'drawing_code', 'row_id', 'symbol_tag', 'tag_family',
               'source_text', 'split_from_joined_text', 'x_pt', 'y_pt', 'floor_x_ft', 'floor_y_ft', 'rotation_deg',
               'footprint_pt', 'assembly_id', 'room_method', 'status', 'model_globalid', 'model_name', 'model_class', 'match_status']
CALLOUT_FIELDS = ['callout_id', 'sheet', 'code', 'typ', 'label_box_pt', 'leader_bends', 'leader_path_pt', 'endpoint_pt',
                  'endpoint_floor_ft', 'trace_check', 'target_kind', 'target', 'target_symbol', 'room', 'room_name',
                  'room_method', 'naive_room_by_label_position', 'naive_differs', 'notes']
EXC_FIELDS = ['exception_id', 'severity', 'category', 'sheet', 'room', 'code', 'related_ids', 'description', 'recommended_action', 'site_check']

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--east', required=True); ap.add_argument('--west', required=True)
    ap.add_argument('--out', default='deliverables'); ap.add_argument('--work', default='build')
    ap.add_argument('--skip-stages', action='store_true', help='reuse intermediate files in --work')
    a = ap.parse_args(argv)
    out = os.path.abspath(a.out); work = os.path.abspath(a.work)
    os.makedirs(out, exist_ok=True); os.makedirs(os.path.join(out, 'pilot'), exist_ok=True)
    import pipeline
    if not a.skip_stages:
        pipeline.prepare(work, a.east, a.west)
        pipeline.run_stages(work)
    os.chdir(work)
    import report2, overlay
    o = report2.build_all()
    rows, inst, cos, exc = o['rows'], o['instances'], o['callouts'], o['exceptions']
    sev_order = {'High': 0, 'Review': 1, 'Info': 2}
    exc.sort(key=lambda e: (sev_order.get(e['severity'], 3), e['category'], e['sheet'], e['room']))
    for n, e in enumerate(exc): e['exception_id'] = 'X%03d' % (n + 1)
    # overlays must be drawn after final exception numbering
    write_csv(os.path.join(out, '01_room_code_takeoff.csv'), rows, ROW_FIELDS)
    write_csv(os.path.join(out, '02_instance_list.csv'), inst, INST_FIELDS)
    write_csv(os.path.join(out, '04_exceptions.csv'), exc, EXC_FIELDS)
    write_csv(os.path.join(out, 'callout_catalog.csv'), cos, CALLOUT_FIELDS)
    rooms = [dict(sheet=('A06.04A (east)' if k.startswith('east') else 'A06.04B (west)'), **{kk: vv for kk, vv in v.items() if kk != 'sheet'})
             for k, v in sorted(o['rooms'].items())]
    write_csv(os.path.join(out, 'rooms.csv'), rooms)
    write_csv(os.path.join(out, 'component_tags_by_room.csv'), o['components'])
    # label count vs counted quantity
    lab = Counter(c['code'] for c in cos)
    typ = Counter(c['code'] for c in cos if c['typ'])
    qty = defaultdict(int); units = defaultdict(set)
    for r in rows:
        qty[r['drawing_code']] += int(r['plan_qty']); units[r['drawing_code']].add(r['unit'])
    lvq = [dict(code=c, code_labels_on_sheets=lab[c], labels_marked_TYP=typ[c], rooms_with_code=sum(1 for r in rows if r['drawing_code'] == c),
                plan_quantity_counted=qty[c], unit='; '.join(sorted(units[c])),
                note=('label count equals counted quantity' if lab[c] == qty[c] else 'label count differs from counted quantity'))
           for c in sorted(lab)]
    write_csv(os.path.join(out, 'label_count_vs_quantity.csv'), lvq)
    # model comparison template (model not supplied)
    write_csv(os.path.join(out, 'model_comparison_template.csv'),
              [dict(room_number=r['room_number'], drawing_code=r['drawing_code'], plan_qty=r['plan_qty'], model_qty='', difference='',
                    status='model not supplied', plan_row_ids=r['row_id'], model_globalids='') for r in rows])
    # overlays
    for fn, tag in (('east', 'A06.04A_east'), ('west', 'A06.04B_west')):
        overlay.render(o, fn, os.path.join(out, '03_overlay_%s.png' % tag), os.path.join(out, '03_overlay_%s.pdf' % tag), dpi=200)
    # pilot close-ups
    import numpy as np
    pilots = []
    from report import PILOT
    for fn, room in PILOT:
        S = o['_objs'][fn][0]
        k = S.room_numbers.index(room) + 1
        ys, xs = np.nonzero(S.R == k)
        x0, x1, y0, y1 = xs.min() / 3 - 30, xs.max() / 3 + 30, ys.min() / 3 - 30, ys.max() / 3 + 30
        png = os.path.join(out, 'pilot', 'pilot_%s_%s.png' % (fn, room))
        overlay.render(o, fn, png, None, dpi=300, crop=(x0, y0, x1, y1))
        pilots.append((fn, room, os.path.basename(png)))
    # workbook
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook(); first = True
        for name, data, fields in (('01 room x code', rows, ROW_FIELDS), ('02 instances', inst, INST_FIELDS),
                                   ('04 exceptions', exc, EXC_FIELDS), ('callouts', cos, CALLOUT_FIELDS),
                                   ('rooms', rooms, list(rooms[0].keys())), ('component tags', o['components'], list(o['components'][0].keys())),
                                   ('label vs qty', lvq, list(lvq[0].keys()))):
            ws = wb.active if first else wb.create_sheet(); first = False
            ws.title = name[:31]; ws.append(fields)
            for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDE6F0')
            for r in data:
                ws.append([(', '.join(map(str, r.get(f))) if isinstance(r.get(f), (list, tuple)) else r.get(f)) for f in fields])
            ws.freeze_panes = 'A2'
            for col in ws.columns:
                width = min(60, max(10, max(len(str(c.value or '')) for c in col[:200]) + 2))
                ws.column_dimensions[col[0].column_letter].width = width
        wb.save(os.path.join(out, 'takeoff_workbook.xlsx'))
    except ImportError:
        pass
    # data for summary / html
    json.dump(dict(rows=rows, instances=inst, callouts=cos, exceptions=exc, components=o['components'],
                   label_vs_qty=lvq, rooms=rooms, pilots=pilots), open(os.path.join(out, 'takeoff_data.json'), 'w'), indent=1, default=str)
    import write_docs
    write_docs.main(out)
    print('rows %d, instances %d, callouts %d, exceptions %d' % (len(rows), len(inst), len(cos), len(exc)))
    return 0

if __name__ == '__main__':
    sys.exit(main())
